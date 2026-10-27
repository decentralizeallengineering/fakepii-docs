"""XLSX: merged cells, CONCAT/& formulas that build names (no cached value),
cross-sheet references, rich-text runs, a sheet named after a person, cell
comments with authors, mailto hyperlinks, print header/footer, chart title,
workbook metadata."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment

from ..people import Plan, split_points
from ..truth import TruthBuilder
from .common import FIXED_TIME, normalize_zip



def _sheet_part(i: int) -> str:
    return f"xl/worksheets/sheet{i}.xml"


# openpyxl writes string cells (plain and rich) as inline strings, so their
# bytes live in the sheet part rather than in xl/sharedStrings.xml.
SHEET1, SHEET2 = _sheet_part(1), _sheet_part(2)


def generate_xlsx(plan: Plan, out_path: str | Path) -> TruthBuilder:
    a, b, c = plan.people[:3]
    rng = plan.rng
    tb = TruthBuilder("xlsx", plan.seed, plan.people)
    wb = Workbook()

    props = wb.properties
    props.creator, props.lastModifiedBy = c.full, a.full
    props.title, props.keywords = f"Roster owner {b.username}", b.email
    props.created = props.modified = FIXED_TIME.replace(tzinfo=None)
    core = "docProps/core.xml"
    tb.pii(c.full, "person_name", c, core, "core/creator", visible=False, techniques=["metadata"])
    tb.pii(a.full, "person_name", a, core, "core/lastModifiedBy", visible=False, techniques=["metadata"])
    tb.pii(b.username, "username", b, core, "core/title", visible=False, techniques=["metadata", "glued"])
    tb.pii(b.email, "email", b, core, "core/keywords", visible=False, techniques=["metadata"])

    # ------------------------------------------------------------ sheet 1
    ws = wb.active
    ws.title = "Roster"
    S1 = "sheet[1]"
    ws["A1"] = f"Team roster prepared by {a.full}"
    ws.merge_cells("A1:H1")
    tb.pii(a.full, "person_name", a, SHEET1, f"{S1}!A1", techniques=["merged_cell"])

    headers = ("First", "Last", "Full name", "Email", "Phone", "SSN", "Address", "Username")
    for col, h in enumerate(headers, 1):
        ws.cell(row=3, column=col, value=h)
    for row, person in enumerate((a, b, c), 4):
        ws[f"A{row}"], ws[f"B{row}"] = person.first, person.last
        ws[f"C{row}"] = f'=_xlfn.CONCAT(A{row}," ",B{row})'
        ws[f"D{row}"] = person.email
        ws[f"D{row}"].hyperlink = f"mailto:{person.email}"
        ws[f"E{row}"], ws[f"F{row}"] = person.phone, person.ssn
        ws[f"G{row}"], ws[f"H{row}"] = person.address, person.username
        tb.pii(person.first, "first_name", person, SHEET1, f"{S1}!A{row}", techniques=["table"])
        tb.pii(person.last, "last_name", person, SHEET1, f"{S1}!B{row}", techniques=["table"])
        tb.pii(person.full, "person_name", person, _sheet_part(1), f"{S1}!C{row}",
               techniques=["formula_concat", "computed_value"])
        tb.pii(person.email, "email", person, SHEET1, f"{S1}!D{row}", techniques=["table"])
        tb.pii(person.email, "email", person, "xl/worksheets/_rels/sheet1.xml.rels", f"{S1}!D{row}@hyperlink",
               visible=False, techniques=["hyperlink_target"])
        tb.pii(person.phone, "phone", person, SHEET1, f"{S1}!E{row}", techniques=["table"])
        tb.pii(person.ssn, "national_id", person, SHEET1, f"{S1}!F{row}", techniques=["table"])
        tb.pii(person.address, "postal_address", person, SHEET1, f"{S1}!G{row}", techniques=["table"])
        tb.pii(person.username, "username", person, SHEET1, f"{S1}!H{row}", techniques=["table", "glued"])

    # Name assembled from string literals: never stored whole anywhere.
    k1 = rng.randrange(1, len(b.first))
    k2 = rng.randrange(len(b.first) + 2, len(b.full))
    frags = (b.full[:k1], b.full[k1:k2], b.full[k2:])
    ws["A8"] = "Signed off by"
    ws["C8"] = "=CONCATENATE(" + ",".join(f'"{f}"' for f in frags) + ")"
    tb.pii(b.full, "person_name", b, _sheet_part(1), f"{S1}!C8",
           techniques=["formula_concat_literals", "computed_value"])
    # Name assembled with & from first/last literals.
    ws["D8"] = f'="{c.first}"&" "&"{c.last}"'
    tb.pii(c.full, "person_name", c, _sheet_part(1), f"{S1}!D8", techniques=["formula_ampersand", "computed_value"])
    tb.pii(c.first, "first_name", c, _sheet_part(1), f"{S1}!D8@formula", visible=False, techniques=["formula_source"])
    tb.pii(c.last, "last_name", c, _sheet_part(1), f"{S1}!D8@formula", visible=False, techniques=["formula_source"])

    # Rich text: one cell, three differently formatted runs.
    pieces = split_points(c.full, rng, 3)
    fonts = (InlineFont(b=True), InlineFont(i=True), InlineFont(color="FF1F4E79", u="single"))
    ws["A10"] = CellRichText(*(TextBlock(f, p) for f, p in zip(fonts, pieces)))
    tb.pii(c.full, "person_name", c, SHEET1, f"{S1}!A10", techniques=["split_runs", "rich_text"])

    ws["A11"] = f"Contact: {c.username}/{c.email_alt},"
    tb.pii(c.username, "username", c, SHEET1, f"{S1}!A11", techniques=["glued"])
    tb.pii(c.email_alt, "email", c, SHEET1, f"{S1}!A11", techniques=["glued", "trailing_punctuation"])

    ws["E4"].comment = Comment(f"Verified by {b.full}; alt {b.phone_alt}", c.full)
    comments = "xl/comments/comment1.xml"
    tb.pii(b.full, "person_name", b, comments, f"{S1}!E4@comment", techniques=["comment"])
    tb.pii(b.phone_alt, "phone", b, comments, f"{S1}!E4@comment", techniques=["comment"])
    tb.pii(c.full, "person_name", c, comments, f"{S1}!E4@comment_author", visible=False,
           techniques=["comment_author"])

    ws.oddHeader.center.text = f"Confidential - {a.full}"
    ws.oddFooter.right.text = a.email_alt
    tb.pii(a.full, "person_name", a, _sheet_part(1), f"{S1}/oddHeader", techniques=["header"])
    tb.pii(a.email_alt, "email", a, _sheet_part(1), f"{S1}/oddFooter", techniques=["footer"])

    town, sku, order = (plan.near(k) for k in
                        ("place_named_like_surname", "product_code_phone_shape", "reference_number_ssn_shape"))
    for col, nm in zip("ABC", (town, sku, order)):
        ws[f"{col}13"] = nm.context
        tb.near(nm, SHEET1, f"{S1}!{col}13", techniques=["near_miss"])

    # ------------------------------------------------------------ sheet 2: named after a person
    ws2 = wb.create_sheet(b.full)
    S2 = "sheet[2]"
    tb.pii(b.full, "person_name", b, "xl/workbook.xml", f"{S2}@name", techniques=["sheet_name"])
    for row, (label, value, type_) in enumerate((("Date of birth", b.dob, "date_of_birth"),
                                                 ("SSN", b.ssn, "national_id"),
                                                 ("Street", b.street, "street_address")), 1):
        ws2[f"A{row}"], ws2[f"B{row}"] = label, value
        tb.pii(value, type_, b, SHEET2, f"{S2}!B{row}")
    ws2["B4"] = '=Roster!A5&" "&Roster!B5'
    tb.pii(b.full, "person_name", b, _sheet_part(2), f"{S2}!B4",
           techniques=["formula_ampersand", "cross_sheet_reference", "computed_value"])

    # ------------------------------------------------------------ sheet 3: chart
    ws3 = wb.create_sheet("Hours")
    ws3.append(("Week", "Hours"))
    for week in range(1, 6):
        ws3.append((f"W{week}", rng.randrange(20, 45)))
    chart = BarChart()
    chart.title = f"Hours logged by {a.full}"
    chart.add_data(Reference(ws3, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(ws3, min_col=1, min_row=2, max_row=6))
    ws3.add_chart(chart, "D2")
    tb.pii(a.full, "person_name", a, "xl/charts/chart1.xml", "title[1]", techniques=["chart_title"])

    out_path = Path(out_path)
    wb.save(out_path)
    normalize_zip(out_path)
    return tb
