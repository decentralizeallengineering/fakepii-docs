"""DOCX: split runs, header/footer, comment, footnote, tables with merged
cells, hyperlink, alt text, legacy form field, content control, tracked
deletion, zero-width and NBSP joins, core metadata."""

from __future__ import annotations

import io
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor

from ..people import Plan, split_points
from ..truth import TruthBuilder
from .common import FIXED_ISO, FIXED_TIME, normalize_zip, png_bytes

DOCUMENT = "word/document.xml"
DOC_RELS = "word/_rels/document.xml.rels"


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def docx_path(el) -> str:
    """Location path of a w:p element, in the extractor's grammar."""
    parts: list[str] = []
    node = el
    while True:
        parent = node.getparent()
        tag = _local(node.tag)
        if tag == "body":
            return "body/" + "/".join(reversed(parts))
        if tag in ("hdr", "ftr"):
            return "/".join(reversed(parts))
        if tag in ("comment", "footnote", "endnote"):
            return f"{tag}[{node.get(qn('w:id'))}]/" + "/".join(reversed(parts))
        if tag in ("p", "tbl", "sdt", "tr", "tc"):
            siblings = [s for s in parent if s.tag == node.tag]
            parts.append(f"{tag}[{siblings.index(node) + 1}]")
        node = parent


def _part_of(doc, reltype: str) -> str:
    for rel in doc.part.rels.values():
        if rel.reltype == reltype and not rel.is_external:
            return str(rel.target_part.partname).lstrip("/")
    raise LookupError(reltype)


def _hyperlink(paragraph, url: str, text: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    link.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr>'
        f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'))
    paragraph._p.append(link)


def _formtext(paragraph, name: str, value: str) -> None:
    w = nsdecls("w")
    for xml in (
        f'<w:r {w}><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="{name}"/><w:enabled/>'
        f'<w:calcOnExit w:val="0"/><w:textInput><w:default w:val="{escape(value, {chr(34): "&quot;"})}"/>'
        f'</w:textInput></w:ffData></w:fldChar></w:r>',
        f'<w:r {w}><w:instrText xml:space="preserve"> FORMTEXT </w:instrText></w:r>',
        f'<w:r {w}><w:fldChar w:fldCharType="separate"/></w:r>',
        f'<w:r {w}><w:t xml:space="preserve">{escape(value)}</w:t></w:r>',
        f'<w:r {w}><w:fldChar w:fldCharType="end"/></w:r>',
    ):
        paragraph._p.append(parse_xml(xml))


def _content_control(paragraph, alias: str, tag: str, value: str) -> None:
    paragraph._p.append(parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val="{alias}"/><w:tag w:val="{tag}"/><w:text/></w:sdtPr>'
        f'<w:sdtContent><w:r><w:t xml:space="preserve">{escape(value)}</w:t></w:r></w:sdtContent></w:sdt>'))


def _tracked_replacement(paragraph, author: str, deleted: str, inserted: str) -> None:
    a = escape(author, {'"': "&quot;"})
    paragraph._p.append(parse_xml(
        f'<w:del {nsdecls("w")} w:id="101" w:author="{a}" w:date="{FIXED_ISO}"><w:r>'
        f'<w:delText xml:space="preserve">{escape(deleted)}</w:delText></w:r></w:del>'))
    paragraph._p.append(parse_xml(
        f'<w:ins {nsdecls("w")} w:id="102" w:author="{a}" w:date="{FIXED_ISO}"><w:r>'
        f'<w:t xml:space="preserve">{escape(inserted)}</w:t></w:r></w:ins>'))


def _add_footnotes(doc, text: str) -> str:
    w = nsdecls("w")
    xml = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes {w}>'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r>'
        '</w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/>'
        f'</w:r><w:r><w:t xml:space="preserve"> {escape(text)}</w:t></w:r></w:p></w:footnote></w:footnotes>'
    )
    part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, xml.encode("utf-8"), doc.part.package)
    doc.part.relate_to(part, RT.FOOTNOTES)
    return "word/footnotes.xml"


def generate_docx(plan: Plan, out_path: str | Path) -> TruthBuilder:
    a, b, c = plan.people[:3]
    rng = plan.rng
    town, sku, order = (plan.near(k) for k in
                        ("place_named_like_surname", "product_code_phone_shape", "reference_number_ssn_shape"))
    tb = TruthBuilder("docx", plan.seed, plan.people)
    doc = Document()

    # --- metadata
    cp = doc.core_properties
    meta = {
        "creator": (a.full, a, "person_name"),
        "title": (f"Personnel summary: {b.full}", b, "person_name"),
        "lastModifiedBy": (c.full, c, "person_name"),
        "subject": (f"HR file {a.username}", a, "username"),
        "keywords": (f"{a.email}; roster", a, "email"),
        "description": (f"Reviewed by {b.full}", b, "person_name"),
    }
    cp.author, cp.title, cp.last_modified_by = meta["creator"][0], meta["title"][0], meta["lastModifiedBy"][0]
    cp.subject, cp.keywords, cp.comments = meta["subject"][0], meta["keywords"][0], meta["description"][0]
    cp.created = cp.modified = FIXED_TIME.replace(tzinfo=None)
    cp.revision = 1
    pii_values = {"person_name": "full", "username": "username", "email": "email"}
    for field, (_, person, type_) in meta.items():
        value = getattr(person, pii_values[type_])
        tb.pii(value, type_, person, "docProps/core.xml", f"core/{field}", visible=False, techniques=["metadata"])

    # --- header / footer
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.text = f"Confidential — prepared for {a.full}"
    header_part = str(section.header.part.partname).lstrip("/")
    tb.pii(a.full, "person_name", a, header_part, docx_path(hp._p), techniques=["header"])
    fp = section.footer.paragraphs[0]
    fp.text = f"Questions: {b.email} | {b.phone}"
    footer_part = str(section.footer.part.partname).lstrip("/")
    tb.pii(b.email, "email", b, footer_part, docx_path(fp._p), techniques=["footer"])
    tb.pii(b.phone, "phone", b, footer_part, docx_path(fp._p), techniques=["footer"])

    doc.add_heading("Staff record", level=1)
    doc.add_paragraph(plan.rng.choice(("Please review the schedule below before the Friday planning meeting.",
                                       "All figures are provisional until the quarterly reconciliation closes.")))

    # --- name split across runs with different formatting
    p = doc.add_paragraph("Employee: ")
    for i, piece in enumerate(split_points(a.full, rng, 3)):
        run = p.add_run(piece)
        run.bold, run.italic, run.underline = i == 0, i == 1, i == 2
        if i == 2:
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            run.font.size = Pt(12)
    p.add_run(".")
    tb.pii(a.full, "person_name", a, DOCUMENT, docx_path(p._p), techniques=["split_runs", "formatting"])

    # --- glued tokens, parentheses + trailing comma
    p = doc.add_paragraph(f"Login {a.username} ({a.email}), backup line {a.phone_alt}.")
    tb.pii(a.username, "username", a, DOCUMENT, docx_path(p._p), techniques=["glued"])
    tb.pii(a.email, "email", a, DOCUMENT, docx_path(p._p), techniques=["parentheses", "trailing_punctuation"])
    tb.pii(a.phone_alt, "phone", a, DOCUMENT, docx_path(p._p))

    # --- first / last names on their own
    p = doc.add_paragraph(f"Dear {a.first}, please ask Ms. {b.last} to countersign.")
    tb.pii(a.first, "first_name", a, DOCUMENT, docx_path(p._p))
    tb.pii(b.last, "last_name", b, DOCUMENT, docx_path(p._p))

    # --- SSN split across runs + DOB
    p = doc.add_paragraph("SSN on file: ")
    head, tail = a.ssn[:5], a.ssn[5:]
    p.add_run(head).bold = True
    p.add_run(tail)
    p.add_run(f"; DOB {a.dob}.")
    tb.pii(a.ssn, "national_id", a, DOCUMENT, docx_path(p._p), techniques=["split_runs"])
    tb.pii(a.dob, "date_of_birth", a, DOCUMENT, docx_path(p._p))

    # --- zero-width space inside a name, NBSP between names
    k = rng.randrange(1, len(b.first))
    p = doc.add_paragraph(f"Emergency contacts: {b.first[:k]}​{b.first[k:]} {b.last} and "
                          f"{c.first} {c.last}.")
    tb.pii(b.full, "person_name", b, DOCUMENT, docx_path(p._p), techniques=["zero_width_space"])
    tb.pii(c.full, "person_name", c, DOCUMENT, docx_path(p._p), techniques=["nbsp"])

    # --- hyperlink (visible text + mailto target in the .rels part)
    p = doc.add_paragraph("Write to ")
    _hyperlink(p, f"mailto:{c.email}", c.email)
    p.add_run(" for payroll questions.")
    tb.pii(c.email, "email", c, DOCUMENT, docx_path(p._p), techniques=["hyperlink_text"])
    tb.pii(c.email, "email", c, DOC_RELS, docx_path(p._p) + "/hyperlink[1]@target",
           visible=False, techniques=["hyperlink_target"])

    # --- comment with an author
    p = doc.add_paragraph()
    run = p.add_run("Badge photo approved.")
    comment = doc.add_comment(run, text=f"Checked against {a.full}'s file; call {a.phone}.",
                              author=b.full, initials="")
    comment._comment_elm.set(qn("w:date"), FIXED_ISO)  # python-docx stamps "now"
    comments_part = _part_of(doc, RT.COMMENTS)
    cpath = f"comment[{comment.comment_id}]"
    tb.pii(b.full, "person_name", b, comments_part, f"{cpath}@author", visible=False,
           techniques=["comment_author"])
    tb.pii(a.full, "person_name", a, comments_part, f"{cpath}/p[1]", techniques=["comment"])
    tb.pii(a.phone, "phone", a, comments_part, f"{cpath}/p[1]", techniques=["comment"])

    # --- footnote
    p = doc.add_paragraph("Address verified by the front office.")
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
                          '<w:footnoteReference w:id="1"/></w:r>'))
    footnotes = _add_footnotes(doc, f"Home address: {a.address}. Mailing: {b.street}.")
    tb.pii(a.address, "postal_address", a, footnotes, "footnote[1]/p[1]", techniques=["footnote"])
    tb.pii(b.street, "street_address", b, footnotes, "footnote[1]/p[1]", techniques=["footnote"])

    # --- picture with alt text
    p = doc.add_paragraph()
    shape = p.add_run().add_picture(io.BytesIO(png_bytes()))
    shape._inline.docPr.set("descr", f"Badge photo of {c.full}")
    shape._inline.docPr.set("title", c.email_alt)
    tb.pii(c.full, "person_name", c, DOCUMENT, docx_path(p._p) + "/drawing[1]@descr",
           visible=False, techniques=["alt_text"])
    tb.pii(c.email_alt, "email", c, DOCUMENT, docx_path(p._p) + "/drawing[1]@title",
           visible=False, techniques=["alt_text"])

    # --- legacy form field + content control
    p = doc.add_paragraph("Name on badge: ")
    _formtext(p, "BadgeName", b.full)
    p.add_run("  Direct line: ")
    _content_control(p, "Direct line", "direct_line", b.phone_alt)
    tb.pii(b.full, "person_name", b, DOCUMENT, docx_path(p._p), techniques=["form_field_value"])
    tb.pii(b.full, "person_name", b, DOCUMENT, docx_path(p._p) + "/ffdata[1]@default",
           visible=False, techniques=["form_field_default"])
    tb.pii(b.phone_alt, "phone", b, DOCUMENT, docx_path(p._p), techniques=["content_control"])

    # --- tracked deletion (hidden text + revision author)
    p = doc.add_paragraph("Previous manager: ")
    _tracked_replacement(p, c.full, a.full, "(name withheld)")
    path = docx_path(p._p)
    tb.pii(a.full, "person_name", a, DOCUMENT, f"{path}/deleted", visible=False, techniques=["tracked_deletion"])
    tb.pii(c.full, "person_name", c, DOCUMENT, f"{path}/del[1]@author", visible=False,
           techniques=["revision_author"])
    tb.pii(c.full, "person_name", c, DOCUMENT, f"{path}/ins[1]@author", visible=False,
           techniques=["revision_author"])

    # --- table with horizontally and vertically merged cells
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    top = table.cell(0, 0).merge(table.cell(0, 2))
    top.text = f"Household record for {b.full}"
    tb.pii(b.full, "person_name", b, DOCUMENT, docx_path(top.paragraphs[0]._p),
           techniques=["table", "merged_cell"])
    for col, (value, type_) in enumerate(((c.full, "person_name"), (c.email_alt, "email"), (c.phone, "phone"))):
        cell = table.cell(1, col)
        cell.text = value
        tb.pii(value, type_, c, DOCUMENT, docx_path(cell.paragraphs[0]._p), techniques=["table"])
    tall = table.cell(2, 0).merge(table.cell(3, 0))
    tall.text = f"Mailing address: {c.address}"
    tb.pii(c.address, "postal_address", c, DOCUMENT, docx_path(tall.paragraphs[0]._p),
           techniques=["table", "merged_cell"])
    table.cell(2, 1).text = "DOB"
    table.cell(2, 2).text = c.dob
    tb.pii(c.dob, "date_of_birth", c, DOCUMENT, docx_path(table.cell(2, 2).paragraphs[0]._p), techniques=["table"])
    table.cell(3, 1).text = "Username"
    table.cell(3, 2).text = c.username
    tb.pii(c.username, "username", c, DOCUMENT, docx_path(table.cell(3, 2).paragraphs[0]._p),
           techniques=["table", "glued"])

    # --- near misses
    p = doc.add_paragraph(f"The depot moved to {town.context} last spring; reorder {sku.context} "
                          f"and chase {order.context}.")
    for nm in (town, sku, order):
        tb.near(nm, DOCUMENT, docx_path(p._p), techniques=["near_miss"])
    doc.add_paragraph("This document is synthetic test data generated for masking evaluation.")

    out_path = Path(out_path)
    doc.save(out_path)
    normalize_zip(out_path)
    return tb
