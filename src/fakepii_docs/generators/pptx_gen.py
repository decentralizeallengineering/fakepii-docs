"""PPTX: split runs, speaker notes, table with merged cells, chart title,
picture alt text, run hyperlinks, group shapes, glued tokens, metadata."""

from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from ..people import Plan, split_points
from ..truth import TruthBuilder
from .common import FIXED_TIME, normalize_zip, png_bytes


def _partname(part) -> str:
    return str(part.partname).lstrip("/")


def _index(shapes, shape) -> int:
    return [s.shape_id for s in shapes].index(shape.shape_id) + 1


def generate_pptx(plan: Plan, out_path: str | Path) -> TruthBuilder:
    a, b, c = plan.people[:3]
    rng = plan.rng
    tb = TruthBuilder("pptx", plan.seed, plan.people)
    prs = Presentation()

    cp = prs.core_properties
    cp.author, cp.title, cp.last_modified_by = a.full, f"Onboarding for {b.full}", c.full
    cp.created = cp.modified = FIXED_TIME.replace(tzinfo=None)
    cp.revision = 1
    core = "docProps/core.xml"
    tb.pii(a.full, "person_name", a, core, "core/creator", visible=False, techniques=["metadata"])
    tb.pii(b.full, "person_name", b, core, "core/title", visible=False, techniques=["metadata"])
    tb.pii(c.full, "person_name", c, core, "core/lastModifiedBy", visible=False, techniques=["metadata"])

    def notes(slide, text: str) -> tuple[str, str]:
        ns = slide.notes_slide
        ns.notes_text_frame.text = text
        return _partname(ns.part), f"shape[{_index(ns.shapes, ns.notes_placeholder)}]/p[1]"

    # ------------------------------------------------------------ slide 1: split runs in subtitle
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    part1 = _partname(s1.part)
    s1.shapes.title.text = "Team onboarding"
    subtitle = s1.placeholders[1]
    para = subtitle.text_frame.paragraphs[0]
    para.add_run().text = "Presented by "
    for i, piece in enumerate(split_points(a.full, rng, 3)):
        run = para.add_run()
        run.text = piece
        run.font.bold, run.font.italic = i == 0, i == 1
        if i == 2:
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            run.font.size = Pt(30)
    tb.pii(a.full, "person_name", a, part1, f"shape[{_index(s1.shapes, subtitle)}]/p[1]",
           techniques=["split_runs", "formatting"])
    npart, npath = notes(s1, f"Speaker notes: call {a.first} at {a.phone} or email {a.email}.")
    tb.pii(a.first, "first_name", a, npart, npath, techniques=["speaker_notes"])
    tb.pii(a.phone, "phone", a, npart, npath, techniques=["speaker_notes"])
    tb.pii(a.email, "email", a, npart, npath, techniques=["speaker_notes"])

    # ------------------------------------------------------------ slide 2: table with merged cells
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    part2 = _partname(s2.part)
    s2.shapes.title.text = f"Contacts for {b.full}"
    tb.pii(b.full, "person_name", b, part2, f"shape[{_index(s2.shapes, s2.shapes.title)}]/p[1]",
           techniques=["slide_title"])
    frame = s2.shapes.add_table(4, 3, Inches(0.5), Inches(1.6), Inches(9), Inches(2.4))
    tpath = f"shape[{_index(s2.shapes, frame)}]/tbl"
    table = frame.table
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(0, 0).text = f"Primary contact: {c.full}"
    tb.pii(c.full, "person_name", c, part2, f"{tpath}/tr[1]/tc[1]/p[1]", techniques=["table", "merged_cell"])
    rows = (
        ((c.email, "email", c), (c.phone, "phone", c), (c.ssn, "national_id", c)),
        ((b.username, "username", b), (b.email_alt, "email", b), (b.dob, "date_of_birth", b)),
    )
    for r, row in enumerate(rows, 1):
        for col, (value, type_, person) in enumerate(row):
            table.cell(r, col).text = value
            tb.pii(value, type_, person, part2, f"{tpath}/tr[{r + 1}]/tc[{col + 1}]/p[1]", techniques=["table"])
    for col, kind in enumerate(("place_named_like_surname", "product_code_phone_shape",
                                "reference_number_ssn_shape")):
        nm = plan.near(kind)
        table.cell(3, col).text = nm.context
        tb.near(nm, part2, f"{tpath}/tr[4]/tc[{col + 1}]/p[1]", techniques=["near_miss", "table"])
    npart, npath = notes(s2, f"Remind {b.last} to update {b.address}.")
    tb.pii(b.last, "last_name", b, npart, npath, techniques=["speaker_notes"])
    tb.pii(b.address, "postal_address", b, npart, npath, techniques=["speaker_notes"])

    # ------------------------------------------------------------ slide 3: chart title
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Hours by week"
    data = CategoryChartData()
    data.categories = [f"W{i}" for i in range(1, 6)]
    data.add_series("Hours", [rng.randrange(20, 45) for _ in range(5)])
    gframe = s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(8), Inches(5), data)
    chart = gframe.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = f"Hours logged by {c.full}"
    tb.pii(c.full, "person_name", c, _partname(gframe.chart_part), "title[1]", techniques=["chart_title"])

    # ------------------------------------------------------------ slide 4: glued, hyperlink, alt text, group
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    part4 = _partname(s4.part)
    box = s4.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1.5))
    tf = box.text_frame
    tf.text = f"Handle: {c.username} ({c.email}), ext. {c.phone_alt}"
    bpath = f"shape[{_index(s4.shapes, box)}]"
    tb.pii(c.username, "username", c, part4, f"{bpath}/p[1]", techniques=["glued"])
    tb.pii(c.email, "email", c, part4, f"{bpath}/p[1]", techniques=["parentheses", "trailing_punctuation"])
    tb.pii(c.phone_alt, "phone", c, part4, f"{bpath}/p[1]")
    p2 = tf.add_paragraph()
    p2.add_run().text = "Email "
    link = p2.add_run()
    link.text = b.email
    link.hyperlink.address = f"mailto:{b.email}"
    tb.pii(b.email, "email", b, part4, f"{bpath}/p[2]", techniques=["hyperlink_text"])
    tb.pii(b.email, "email", b, f"ppt/slides/_rels/{part4.rsplit('/', 1)[1]}.rels", f"{bpath}/p[2]/hlink[1]@target",
           visible=False, techniques=["hyperlink_target"])

    pic = s4.shapes.add_picture(io.BytesIO(png_bytes()), Inches(0.5), Inches(2.5), Inches(1.5), Inches(1.5))
    pic._element.nvPicPr.cNvPr.set("descr", f"Headshot of {b.full}")
    pic._element.nvPicPr.cNvPr.set("title", b.username)
    ppath = f"shape[{_index(s4.shapes, pic)}]"
    tb.pii(b.full, "person_name", b, part4, f"{ppath}@descr", visible=False, techniques=["alt_text"])
    tb.pii(b.username, "username", b, part4, f"{ppath}@title", visible=False, techniques=["alt_text"])

    group = s4.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(3), Inches(2.5), Inches(5), Inches(1))
    inner.text_frame.text = f"Badge holder: {a.first} {a.last}, SSN {a.ssn}"
    gpath = f"shape[{_index(s4.shapes, group)}]/shape[{_index(group.shapes, inner)}]/p[1]"
    tb.pii(a.full, "person_name", a, part4, gpath, techniques=["group_shape", "nbsp"])
    tb.pii(a.ssn, "national_id", a, part4, gpath, techniques=["group_shape"])

    out_path = Path(out_path)
    prs.save(out_path)
    normalize_zip(out_path)
    return tb
