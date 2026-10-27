"""Every format is generated, reproducibly, with fake-only data and the
required hard places."""

import re

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from fakepii_docs import FORMATS, data, generate_corpus
from fakepii_docs.extract import raw_texts
from fakepii_docs.truth import sha256_file, validate

from .conftest import SEED


def test_every_format_has_document_and_valid_truth(corpus):
    for fmt in FORMATS:
        entry = corpus[fmt]
        assert entry["doc"].stat().st_size > 0
        validate(entry["data"])
        assert entry["data"]["document"]["format"] == fmt
        assert entry["data"]["document"]["sha256"] == sha256_file(entry["doc"])
        assert entry["data"]["seed"] == SEED
        assert len(entry["data"]["pii"]) >= 10
        assert {e["entity"] for e in entry["data"]["pii"]} == {"person-1", "person-2", "person-3"}


def test_same_seed_same_bytes(corpus, tmp_path):
    generate_corpus(tmp_path, SEED)
    for fmt in FORMATS:
        doc, truth = corpus[fmt]["doc"], corpus[fmt]["truth"]
        assert sha256_file(tmp_path / doc.name) == sha256_file(doc), fmt
        assert sha256_file(tmp_path / truth.name) == sha256_file(truth), fmt


def test_other_seed_other_people(corpus, tmp_path):
    generate_corpus(tmp_path, SEED + 1, ["docx"])
    other = (tmp_path / f"fakepii-{SEED + 1}-docx.docx.truth.json").read_text("utf-8")
    assert sha256_file(tmp_path / f"fakepii-{SEED + 1}-docx.docx") != sha256_file(corpus["docx"]["doc"])
    assert other != corpus["docx"]["truth"].read_text("utf-8")


PATTERNS = {
    "email": r"[a-z.]+@example\.(invalid|com)",
    "phone": r"(\(\d{3}\) 555-01\d\d|\+1 \d{3} 555 01\d\d)",
    "national_id": r"(9\d\d|000)-\d\d-\d{4}",
    "postal_address": r"\d+ [A-Za-z ]+, [A-Za-z]+, ZZ 000\d\d",
    "street_address": r"\d+ [A-Za-z ]+",
    "date_of_birth": r"\d{4}-\d\d-\d\d",
}


def test_values_are_fake_only(corpus):
    for fmt in FORMATS:
        for entry in corpus[fmt]["data"]["pii"]:
            value, type_ = entry["value"], entry["type"]
            if type_ in PATTERNS:
                assert re.fullmatch(PATTERNS[type_], value), (fmt, type_, value)
                if type_ in ("postal_address", "street_address"):
                    assert any(street in value for street in data.STREETS)
            elif type_ == "first_name":
                assert value in data.FIRST_NAMES
            elif type_ == "last_name":
                assert value in data.LAST_NAMES
            elif type_ == "person_name":
                first, last = value.split(" ")
                assert first in data.FIRST_NAMES and last in data.LAST_NAMES
            elif type_ == "username":
                assert re.fullmatch(r"[A-Z][a-z]+[A-Z][a-z]+(19|20)\d\d", value)


def test_name_pools_do_not_collide_with_templates_or_filler(tmp_path):
    Document().save(tmp_path / "blank.docx")
    Presentation().save(tmp_path / "blank.pptx")
    Workbook().save(tmp_path / "blank.xlsx")
    words = data.FIRST_NAMES + data.LAST_NAMES
    texts = [t for name in ("blank.docx", "blank.pptx", "blank.xlsx") for _, t in raw_texts(tmp_path / name)]
    texts += list(data.FILLER)
    for word in words:
        assert not any(word in t for t in texts), word
    for a in words:
        for b in words:
            assert a == b or a not in b, (a, b)


REQUIRED_TECHNIQUES = {
    "split_runs", "glued", "parentheses", "trailing_punctuation", "header", "footer", "comment",
    "comment_author", "footnote", "table", "merged_cell", "formula_concat", "formula_concat_literals",
    "sheet_name", "chart_title", "speaker_notes", "form_field_value", "metadata", "hyperlink_target",
    "alt_text", "tracked_deletion", "zero_width_space", "nbsp", "line_wrap", "embedded_newline",
    "split_emphasis", "front_matter", "html_comment",
}


def test_hard_places_are_covered(corpus):
    seen = {t for fmt in FORMATS for e in corpus[fmt]["data"]["pii"] for loc in e["locations"]
            for t in loc["techniques"]}
    assert REQUIRED_TECHNIQUES <= seen, REQUIRED_TECHNIQUES - seen
    hidden = {fmt for fmt in FORMATS for e in corpus[fmt]["data"]["pii"] for loc in e["locations"]
              if not loc["visible"]}
    assert hidden == set(FORMATS) - {"csv", "txt"}


@pytest.mark.parametrize("fmt", FORMATS)
def test_labeled_near_misses_in_every_document(corpus, fmt):
    near = corpus[fmt]["data"]["near_misses"]
    assert {n["kind"] for n in near} == {"place_named_like_surname", "product_code_phone_shape",
                                         "reference_number_ssn_shape"}
    surnames = {e["display"].split()[1] for e in corpus[fmt]["data"]["entities"]}
    for n in near:
        assert n["label"] and n["value"] in n["context"] and n["locations"]
        if n["kind"] == "place_named_like_surname":
            assert n["value"] in surnames
        if n["kind"] == "product_code_phone_shape":
            assert re.fullmatch(r"\d{3}-555-01\d\d", n["value"])
        values = {e["value"] for e in corpus[fmt]["data"]["pii"]}
        assert n["value"] not in values or n["kind"] == "place_named_like_surname"
